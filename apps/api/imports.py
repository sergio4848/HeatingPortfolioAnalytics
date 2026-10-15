import csv
import io
import json
import zipfile
from pathlib import PurePosixPath
from openpyxl import load_workbook
from .analysis import validate

MAX_BYTES = 2_000_000
MAX_ROWS = 10000


def read_upload(name, content, mapping=None):
    if len(content) > MAX_BYTES:
        raise ValueError("File limit is 2 MB")
    suffix = PurePosixPath(name).suffix.lower()
    if suffix == ".csv":
        text = content.decode("utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        rows = []
        for i, row in enumerate(reader):
            if i >= MAX_ROWS:
                raise ValueError("Row limit is 10,000")
            rows.append(row)
        headers = reader.fieldnames or []
    elif suffix == ".xlsx":
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 300 or sum(x.file_size for x in entries) > 20_000_000:
                raise ValueError("Expanded workbook limit exceeded")
            if any(
                ".." in PurePosixPath(x.filename).parts
                or x.filename.startswith("/")
                or "\\" in x.filename
                or "vbaproject" in x.filename.lower()
                or "externallinks/" in x.filename.lower()
                for x in entries
            ):
                raise ValueError("Unsafe or macro-enabled workbook")
            if "[Content_Types].xml" not in archive.namelist():
                raise ValueError("Invalid XLSX content")
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        sheet = wb.active
        iterator = sheet.iter_rows()
        headers = [str(c.value or "") for c in next(iterator)]
        rows = []
        for i, cells in enumerate(iterator):
            if i >= MAX_ROWS:
                raise ValueError("Row limit is 10,000")
            if any(c.data_type == "f" for c in cells):
                raise ValueError(f"Formula in workbook row {i + 2}; export values only")
            rows.append(dict(zip(headers, [c.value for c in cells])))
        wb.close()
    else:
        raise ValueError("Only CSV and macro-free XLSX supported")
    if len(headers) != len(set(headers)) or not headers or len(headers) > 60:
        raise ValueError("Unique headers required; maximum 60 columns")
    if not rows:
        raise ValueError("File contains no data rows")
    maps = json.loads(mapping) if isinstance(mapping, str) and mapping else mapping or {}
    if not isinstance(maps, dict):
        raise ValueError("Column mapping must be an object")
    mapped = [{maps.get(k, k): v for k, v in row.items()} for row in rows]
    clean, issues = validate(mapped)
    return {
        "headers": headers,
        "preview": mapped[:8],
        "records": clean,
        "issues": issues,
        "total_rows": len(rows),
    }


def safe_cell(value):
    text = str(value)
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")) else text
